"""Contract tests for the mock adapters (same interface the live ones implement)."""

from pipeline.adapters.etsy import get_etsy_client
from pipeline.adapters.images import get_image_provider
from pipeline.adapters.printify import get_printify_client


def test_mock_etsy_contract():
    c = get_etsy_client()
    results = c.search_listings("minimalist botanical print", limit=5)
    assert results and all(r.listing_id and r.title and r.price_usd > 0 for r in results)
    assert c.search_listings("minimalist botanical print", limit=5)[0].title == results[0].title  # deterministic
    nodes = c.flatten_taxonomy()
    assert any("Digital Prints" in n.full_path for n in nodes)
    d = c.create_draft_listing({"title": "Test", "tags": ["a"], "price": 5.0})
    assert d.state == "draft" and d.listing_id
    img = c.upload_listing_image(d.listing_id, b"\x89PNG", "a.png", rank=1)
    f = c.upload_listing_file(d.listing_id, b"%PDF", "a.pdf", "a.pdf")
    assert img["uri"] and f["uri"]
    assert c.get_listing_stats([d.listing_id])[0].state == "draft"
    assert c.get_shop_receipts() == []


def test_mock_etsy_persists_across_instances():
    from pipeline.adapters.etsy.mock import MockEtsyClient

    a = MockEtsyClient()
    d = a.create_draft_listing({"title": "persist"})
    b = MockEtsyClient()
    assert b.get_listing_stats([d.listing_id])


def test_mock_image_provider_is_prompt_dependent():
    from pipeline.agents.qa import compute_phash, hamming

    p = get_image_provider()
    a = p.generate("eucalyptus stem line drawing", size=(512, 640))
    b = p.generate("folded paper crane line drawing", size=(512, 640))
    assert a.media_type == "image/png" and a.width == 512 and a.cost_usd == 0.0
    assert hamming(compute_phash(a.data), compute_phash(b.data)) > 6
    assert a.data == p.generate("eucalyptus stem line drawing", size=(512, 640)).data  # deterministic


def test_mock_printify_contract():
    import pytest

    c = get_printify_client()
    assert c.list_shops() and c.list_blueprints()
    prod = c.create_product("1", {"title": "poster"})
    assert prod["id"]
    with pytest.raises(NotImplementedError):
        c.publish_product("1", prod["id"])


def test_local_storage_roundtrip():
    from pipeline.storage import get_storage

    st = get_storage()
    uri = st.put("products/1/a.txt", b"hello", "text/plain")
    assert uri.startswith("local://")
    assert st.get(uri) == b"hello" and st.exists(uri)
    assert "products/1/a.txt" in st.list("products")
    st.delete(uri)
    assert not st.exists(uri)


def test_generators_produce_valid_files():
    import io

    from openpyxl import load_workbook

    from pipeline.generators.pdf import render_planner_pdf
    from pipeline.generators.xlsx import render_spreadsheet

    pdf = render_planner_pdf({"title": "T", "subtitle": "s", "sections": ["A", "B"], "pages": 2}, kind="tracker")
    assert pdf.startswith(b"%PDF")
    xlsx = render_spreadsheet({"title": "Book", "sheets": [{"name": "Income", "columns": ["Date", "Amount"]}]})
    wb = load_workbook(io.BytesIO(xlsx))
    # default 50 data rows (4-53), totals row 54; "Amount" inferred as currency
    assert wb.sheetnames == ["Income"] and wb["Income"]["B54"].value == "=SUM(B4:B53)"
    typed = render_spreadsheet({"title": "Book", "sheets": [{"name": "Summary", "rows": 5, "totals": False, "columns": [
        {"name": "Income", "type": "currency"}, {"name": "Costs", "type": "currency"},
        {"name": "Profit", "type": "formula", "formula": "=A{r}-B{r}"}]}]})
    ws = load_workbook(io.BytesIO(typed))["Summary"]
    assert ws["C4"].value == "=A4-B4" and ws["C8"].value == "=A8-B8" and ws["A9"].value is None
