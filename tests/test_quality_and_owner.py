"""Premium spreadsheet engine, validation gate, real-file QA and owner feedback."""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from pipeline.generators import office
from pipeline.generators.xlsx import (
    SpreadsheetError,
    render_spreadsheet,
    sanitize_formula,
    spec_problems,
    validate_xlsx,
)

needs_lo = pytest.mark.skipif(not office.available(), reason="LibreOffice not installed")


def test_sanitizer_fixes_what_used_to_corrupt_files():
    m = {"monthly summary": "Monthly Summary", "income": "Income"}
    f, p = sanitize_formula("=A{r}/A{r-1}-1", row=5, sheet_map=m)
    assert f == "=A5/A4-1" and not p
    f, p = sanitize_formula("=SUMIFS(Monthly Summary!B:B, Income!A:A, A{r})", row=5, sheet_map=m)
    assert f == "=SUMIFS('Monthly Summary'!B:B, 'Income'!A:A, A5)" and not p
    f, p = sanitize_formula('=IF(B{r}="Income!x",1,0)', row=5, sheet_map=m)
    assert f == '=IF(B5="Income!x",1,0)'
    assert sanitize_formula("=XLOOKUP(A{r},Income!A:A,Income!B:B)", row=5, sheet_map=m)[1]
    assert sanitize_formula("=SUM(Nope!A:A)", row=5, sheet_map=m)[1]
    assert sanitize_formula("=(A{r}+1", row=5, sheet_map=m)[1]


def test_bad_spec_is_reported_not_shipped():
    spec = {"title": "T", "sheets": [{"name": "Data", "columns": [
        {"name": "A", "type": "number"}, {"name": "B", "type": "formula", "formula": "=FILTER(A:A,A:A>1)"}]}]}
    assert any("FILTER" in p for p in spec_problems(spec))
    with pytest.raises(SpreadsheetError):
        render_spreadsheet(spec, strict=True)


def premium_spec():
    from pipeline.llm import _mock_sheet_spec

    return _mock_sheet_spec("Bookkeeping")


def test_premium_workbook_structure():
    wb = load_workbook(io.BytesIO(render_spreadsheet(premium_spec())))
    assert wb.sheetnames == ["Start Here", "Dashboard", "Income", "Expenses", "Mileage Log", "Settings"]
    dash = wb["Dashboard"]
    assert dash._charts and any(str(c.value).startswith("=SUM(") for row in dash.iter_rows() for c in row)
    income = wb["Income"]
    assert any("list_" in (dv.formula1 or "") for dv in income.data_validations.dataValidation)  # dropdowns
    assert income["A4"].value is not None  # example entries ship in the file
    assert "list_1" in wb.defined_names


@needs_lo
def test_validation_gate_recalculates_with_zero_errors():
    report = validate_xlsx(render_spreadsheet(premium_spec()))
    assert report["ok"] and report["recalculated"] and report["formulas"] > 50


@needs_lo
def test_validation_gate_catches_formula_errors():
    from openpyxl import Workbook

    wb = Workbook()
    wb.active["A1"] = "=1/0"
    buf = io.BytesIO()
    wb.save(buf)
    report = validate_xlsx(buf.getvalue())
    assert not report["ok"] and "#DIV/0!" in report["errors"][0]


def test_unparseable_placeholder_is_caught_without_libreoffice():
    from openpyxl import Workbook

    wb = Workbook()
    wb.active["A1"] = "=A4/A{r-1}-1"
    buf = io.BytesIO()
    wb.save(buf)
    report = validate_xlsx(buf.getvalue(), recalc=False)
    assert not report["ok"] and "placeholder" in report["errors"][0]


HDR = {"X-Siliconopolis": "1"}


def _client():
    from fastapi.testclient import TestClient

    from pipeline.web.app import app

    return TestClient(app)


def _built_product(session, product_type="spreadsheet"):
    from pipeline.agents.creator import CreatorAgent
    from pipeline.db.models import Product

    p = Product(state="brief", niche="sheets", product_type=product_type,
                brief_json={"concept": "Bookkeeping workbook", "target_price_usd": 12})
    session.add(p)
    session.commit()
    CreatorAgent().run()
    session.expire_all()
    return session.get(Product, p.id)


def test_revision_request_sends_product_back_with_feedback(session):
    from pipeline.agents.creator import CreatorAgent
    from pipeline.db.models import OwnerNote, Product, ProductState
    p = _built_product(session)
    c = _client()
    assert c.post(f"/api/dashboard/product/{p.id}/comments", json={"text": "x"}).status_code == 403  # no header
    r = c.post(f"/api/dashboard/product/{p.id}/comments", json={"text": "Love the colours"}, headers=HDR)
    assert r.status_code == 200 and r.json()["state"] == ProductState.CREATED
    r = c.post(f"/api/dashboard/product/{p.id}/comments", headers=HDR,
               json={"text": "Make the dashboard warmer and add a mileage tab", "revise": True})
    assert r.status_code == 200 and r.json()["state"] == ProductState.BRIEF
    detail = c.get(f"/api/dashboard/product/{p.id}").json()
    assert detail["revision_pending"] and len(detail["comments"]) == 2 and not detail["can_revise"]

    CreatorAgent().run()
    session.expire_all()
    p = session.get(Product, p.id)
    assert p.state == ProductState.CREATED and "revision" not in p.brief_json
    assert p.brief_json["revisions_done"][0]["feedback"].count("- ") == 2  # earlier note rides along
    from pipeline.agents.creator import revision_note

    note = revision_note({"revision": {"feedback": "- warmer", "round": 1, "previous_plan": {"title_hint": "Old"}}})
    assert "OWNER REVISION REQUEST" in note and "warmer" in note and "Old" in note
    statuses = {n.kind: n.status for n in session.query(OwnerNote).filter_by(product_id=p.id)}
    assert statuses == {"comment": "included", "revision": "addressed"}


def test_revision_refused_for_etsy_drafts(session):
    from pipeline.db.models import Product

    p = Product(state="drafted", niche="n", product_type="planner", brief_json={})
    session.add(p)
    session.commit()
    r = _client().post(f"/api/dashboard/product/{p.id}/comments", json={"text": "x", "revise": True}, headers=HDR)
    assert r.status_code == 409 and "Etsy" in r.json()["detail"]


def test_run_brief_steers_exactly_one_research_run(session):
    from pipeline.agents.researcher import BRIEF_NICHE, ResearcherAgent
    from pipeline.db.models import BriefStatus, Product, ResearchRequest
    from pipeline.llm import get_llm

    c = _client()
    assert c.post("/api/dashboard/research/briefs", json={"focus": "Halloween decor"}).status_code == 403  # header
    b = c.post("/api/dashboard/research/briefs", headers=HDR,
               json={"focus": "Halloween decor, cute not scary", "product_types": ["mug", "nonsense"], "idea_count": 3}).json()["brief"]
    assert b["status"] == "queued" and b["product_types"] == ["mug"]
    second = c.post("/api/dashboard/research/briefs", headers=HDR, json={"focus": "Wedding planners"}).json()["brief"]
    edited = c.patch(f"/api/dashboard/research/briefs/{second['id']}", headers=HDR,
                     json={"focus": "Wedding seating spreadsheets", "idea_count": 2}).json()["brief"]
    assert edited["focus"] == "Wedding seating spreadsheets"
    assert [q["id"] for q in c.get("/api/dashboard/research").json()["queued"]] == [b["id"], second["id"]]

    ResearcherAgent().run()  # uses ONLY the first brief
    calls = [x for x in get_llm().calls if x["task"] == "research_briefs"]
    assert len(calls) == 1 and "OWNER RUN BRIEF" in calls[0]["user"] and "Halloween" in calls[0]["user"]
    session.expire_all()
    done = session.get(ResearchRequest, b["id"])
    assert done.status == BriefStatus.DONE and len(done.product_ids) == 3
    made = session.query(Product).filter(Product.id.in_([int(i) for i in done.product_ids])).all()
    assert all(p.product_type == "mug" and p.brief_json["from_brief"] == b["id"] for p in made)
    assert session.get(ResearchRequest, second["id"]).status == BriefStatus.QUEUED
    assert c.patch(f"/api/dashboard/research/briefs/{b['id']}", headers=HDR, json={"focus": "x" * 5}).status_code == 409

    research = c.get("/api/dashboard/research").json()
    assert research["runs"][0]["brief"]["id"] == b["id"] and len(research["runs"][0]["ideas"]) == 3
    assert c.delete(f"/api/dashboard/research/briefs/{second['id']}", headers=HDR).json()["brief"]["status"] == "cancelled"

    ResearcherAgent().run()  # no brief left: back to the usual niche list
    session.expire_all()
    assert session.query(Product).filter(Product.niche == BRIEF_NICHE).count() <= 3
    calls = [x for x in get_llm().calls if x["task"] == "research_briefs"]
    assert len(calls) > 2 and not any("OWNER RUN BRIEF" in x["user"] for x in calls[1:])


def test_delete_products_single_bulk_and_all_rejected(session):
    from pipeline.db.models import Product
    from pipeline.storage import get_storage

    storage = get_storage()
    keys = []
    for i, state in enumerate(["rejected", "rejected", "brief", "drafted", "approved"]):
        key = storage.put(f"products/{i}-x/listing-01-primary.png", b"png", "image/png")
        storage.put(f"products/{i}-x/source/design.html", b"<html>", "text/html")
        keys.append(key)
        session.add(Product(state=state, niche="n", product_type="planner", brief_json={"concept": f"c{i}"},
                            image_paths=[key]))
    session.commit()
    ids = [p.id for p in session.query(Product).order_by(Product.id)]
    c = _client()
    assert c.post("/api/dashboard/products/delete", json={"state": "rejected"}).status_code == 403
    r = c.post("/api/dashboard/products/delete", headers=HDR, json={"state": "rejected"}).json()
    assert sorted(r["deleted"]) == ids[:2] and not storage.exists(keys[0])
    assert not storage.list("products/0-x/")  # whole product folder removed
    r = c.post("/api/dashboard/products/delete", headers=HDR, json={"ids": ids[2:]}).json()
    assert r["deleted"] == [ids[2], ids[4]] and r["skipped"][0]["id"] == ids[3]  # Etsy drafts stay
    assert c.delete(f"/api/dashboard/product/{ids[3]}", headers=HDR).status_code == 409
    session.expire_all()
    assert [p.state for p in session.query(Product)] == ["drafted"]


def test_qa_reviews_real_file_and_enforces_commercial_bar(session, monkeypatch):
    from pipeline.agents.qa import QAAgent
    from pipeline.db.models import Product, ProductState
    from pipeline.llm import MockLLM

    p = _built_product(session, "planner")
    orig = MockLLM._canned

    def weak(task, user):
        out = orig(task, user)
        if task == "qa_vision":
            out["scores"]["commercial_viability"] = 6
        return out

    monkeypatch.setattr(MockLLM, "_canned", staticmethod(weak))
    QAAgent().run()
    session.expire_all()
    p = session.get(Product, p.id)
    qa = p.qa_result_json
    assert p.state == ProductState.REJECTED and qa["reviewed"]["file_images"] >= 1
    assert any("commercial_viability 6" in r for r in qa["reasons"])
