"""Premium spreadsheet templates (openpyxl), a validation gate, and HTML/real-file previews.

spec = {
  "title": str, "subtitle": str, "font": str,
  "instructions": [str, ...],                      # "Start Here" tab how-to steps
  "sheets": [{"name": str, "purpose": str, "rows": int, "totals": bool,
              "columns": [str | {"name", "type", "formula", "format", "width", "options", "highlight"}],
              "sample_rows": [[...]]}],
  "dashboard": {"name": str,
                "kpis": [{"label", "agg": sum|average|count|max|min|difference|ratio, "sheet", "column",
                          "where": {"column", "equals"} | None, "of": [label, label], "format"}],
                "breakdowns": [{"title", "sheet", "group_by", "value_column", "agg": sum|count,
                                "chart": column|bar|pie|doughnut|line}]}}

Column types: text, date, currency, number, percent, checkbox, choice (dropdown from `options`),
formula (`formula` uses {r} for this row, {r-1} for the row above).

The workbook gets a Start Here tab, an optional Dashboard (KPI cards, summary tables, native charts),
the data tabs (styled, dropdowns, conditional formatting, print-ready) and a Settings tab holding the
dropdown lists so buyers can rename categories. Example rows are included so charts are alive on open.
Every formula is sanitised (row placeholders, quoted sheet names, blank-row guards) and the result is
validated by recalculating it in LibreOffice: a file with any formula error never leaves the Creator.
"""

from __future__ import annotations

import html as html_lib
import io
import re
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, PieChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from pipeline.generators.mockups import short_title

NUMBER_FORMATS = {"date": "mmm d, yyyy", "currency": '"$"#,##0.00;[Red]-"$"#,##0.00', "number": "#,##0.##",
                  "percent": "0.0%"}
COLUMN_TYPES = (*NUMBER_FORMATS, "text", "checkbox", "choice", "formula")
SUMMABLE = ("currency", "number")
HEADER_ROW = 3
DEFAULT_ROWS = 100
SAFE_FONTS = ("Calibri", "Arial", "Georgia", "Verdana", "Trebuchet MS", "Century Gothic", "Cambria", "Tahoma")
CHECK_ON, CHECK_OFF = "☑", "☐"
SETTINGS = "Settings"
START = "Start Here"
SPARE_OPTIONS = 10
# Functions newer than Excel 2010 break in older Excel / need _xlfn prefixes: never allowed.
BANNED_FUNCS = {"XLOOKUP", "XMATCH", "FILTER", "UNIQUE", "SORT", "SORTBY", "SEQUENCE", "LET", "LAMBDA", "TEXTJOIN",
                "CONCAT", "IFS", "SWITCH", "MAXIFS", "MINIFS", "RANDARRAY", "TAKE", "DROP", "CHOOSECOLS", "CHOOSEROWS",
                "VSTACK", "HSTACK", "TEXTSPLIT", "TEXTBEFORE", "TEXTAFTER", "TOCOL", "TOROW", "WRAPROWS", "IMAGE",
                "GOOGLEFINANCE", "QUERY", "ARRAYFORMULA", "IMPORTRANGE", "SPARKLINE"}
EXCEL_ERRORS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A", "Err:", "#ERR")
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
STATUS_WORDS = ("status", "stage", "priority", "state", "progress")
STATUS_TINTS = ["DDEFE0", "FFF1CC", "FADBD8", "DCE8F7", "EADCF5", "F3E5D8", "E0F2F1", "ECEFF1"]


class SpreadsheetError(ValueError):
    """The spec cannot be built into a clean workbook; `problems` are fed back to the Creator to fix."""

    def __init__(self, problems: list[str]):
        super().__init__("; ".join(problems[:8]))
        self.problems = problems


# ---------------------------------------------------------------------------------------------- spec helpers
MONEY_WORDS = ("amount", "income", "expense", "profit", "total", "cost", "price", "balance", "revenue", "budget", "$",
               "outstanding", "owed", "billed", "fee", "spent", "salary", "payment", "deposit")
PERCENT_WORDS = ("margin", "%", "percent", "rate", "change", "growth", "share", "ratio")


def normalize_columns(cols: list[Any]) -> list[dict[str, Any]]:
    out, seen = [], set()
    for c in cols[:20]:
        if isinstance(c, dict):
            name = str(c.get("name") or "Column").strip() or "Column"
            ctype = str(c.get("type") or "text").lower()
            fmt = str(c.get("format") or "").lower() or None
            options = [str(o).strip() for o in (c.get("options") or []) if str(o).strip()][:30]
            if ctype in ("dropdown", "select", "list", "category"):
                ctype = "choice"
            if ctype == "choice" and not options:
                ctype = "text"
            col = {"name": name, "type": ctype if ctype in COLUMN_TYPES else "text", "formula": c.get("formula"),
                   "format": fmt if fmt in NUMBER_FORMATS else None, "width": c.get("width"), "options": options,
                   "highlight": c.get("highlight")}
        else:
            name = str(c).strip() or "Column"
            lowered = name.lower()
            ctype = "checkbox" if lowered.endswith("?") else "date" if "date" in lowered \
                else "currency" if any(k in lowered for k in MONEY_WORDS) else "text"
            col = {"name": name, "type": ctype, "formula": None, "format": None, "width": None, "options": [],
                   "highlight": None}
        base, n = col["name"], 2
        while col["name"].lower() in seen:  # duplicate headers confuse filters and lookups
            col["name"] = f"{base} {n}"
            n += 1
        seen.add(col["name"].lower())
        out.append(col)
    return out or [{"name": "Column A", "type": "text", "formula": None, "format": None, "width": None,
                    "options": [], "highlight": None}]


def _display_type(col: dict[str, Any]) -> str:
    if col.get("format"):
        return col["format"]
    if col["type"] in ("formula", "number"):
        name = col["name"].lower()
        if any(w in name for w in PERCENT_WORDS):
            return "percent"
        if any(w in name for w in MONEY_WORDS):
            return "currency"
        if col["type"] == "formula":
            return "number"
    return col["type"]


def safe_sheet_name(name: str, used: set[str]) -> str:
    clean = re.sub(r"[\[\]:*?/\\]", " ", str(name or "Sheet")).strip().strip("'").strip() or "Sheet"
    clean = re.sub(r"\s+", " ", clean)[:31].strip()
    if clean.lower() == "history":
        clean = "History Log"
    base, n = clean, 2
    while clean.lower() in used:
        suffix = f" {n}"
        clean = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(clean.lower())
    return clean


def _hex(c: str | None, default: str) -> str:
    c = (c or default).lstrip("#")
    return c.upper() if re.fullmatch(r"[0-9A-Fa-f]{6}", c) else default.lstrip("#").upper()


def _mix(hex_a: str, hex_b: str, t: float) -> str:
    a = [int(hex_a[i:i + 2], 16) for i in (0, 2, 4)]
    b = [int(hex_b[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(a, b, strict=True))


def _luma(hex_c: str) -> float:
    r, g, b = (int(hex_c[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def _qs(name: str) -> str:
    return "'" + name.replace("'", "''") + "'"


def _str_lit(s: str) -> str:
    return '"' + str(s).replace('"', '""') + '"'


# ---------------------------------------------------------------------------------------------- formulas
_PLACEHOLDER = re.compile(r"\{\s*(?:r|row)\s*(?:([+-])\s*(\d+))?\s*\}", re.I)
_FUNC = re.compile(r"(?<![A-Za-z0-9_.])([A-Za-z][A-Za-z0-9_.]*)\s*\(")
_CELL = re.compile(r"(?<![A-Za-z0-9_!$'\"])\$?([A-Z]{1,3})\$?(\d{1,7})(?![A-Za-z0-9_(!])")


def _strip_strings(formula: str) -> str:
    return re.sub(r'"(?:[^"]|"")*"', '""', formula)


def sanitize_formula(formula: str, *, row: int, sheet_map: dict[str, str]) -> tuple[str, list[str]]:
    """Make a model-written formula valid Excel. Returns (formula, problems)."""
    problems: list[str] = []
    f = str(formula or "").strip()
    if not f:
        return "", ["empty formula"]
    if not f.startswith("="):
        f = "=" + f
    f = _PLACEHOLDER.sub(lambda m: str(max(1, row + (int(m.group(2)) * (1 if m.group(1) == "+" else -1)
                                                      if m.group(1) else 0))), f)
    strings: list[str] = []  # string literals are set aside so sheet-name rewriting never touches them

    def keep(m: re.Match) -> str:
        strings.append(m.group(0))
        return f"\x01{len(strings) - 1}\x01"

    f = re.sub(r'"(?:[^"]|"")*"', keep, f)
    # sheet references: quote every known sheet (renamed sheets are mapped to their final names)
    tokens: dict[str, str] = {}

    def protect(final: str) -> str:
        key = f"\x00{len(tokens)}\x00"
        tokens[key] = _qs(final) + "!"
        return key

    def quoted(m: re.Match) -> str:
        name = m.group(1).replace("''", "'")
        final = sheet_map.get(name.lower())
        if final is None:
            problems.append(f"formula references unknown sheet '{name}'")
            return m.group(0)
        return protect(final)

    f = re.sub(r"'((?:[^']|'')+)'!", quoted, f)
    for original in sorted(sheet_map, key=len, reverse=True):
        pattern = re.compile(r"(?<![A-Za-z0-9_.\x00])" + re.escape(original) + r"!", re.I)
        f = pattern.sub(lambda _m, o=original: protect(sheet_map[o]), f)
    leftover = re.findall(r"(?<![A-Za-z0-9_.\x00'])([A-Za-z_][A-Za-z0-9_.]*)!", _strip_strings(f))
    for name in leftover:
        problems.append(f"formula references unknown sheet '{name}'")
    for key, val in tokens.items():
        f = f.replace(key, val)
    f = re.sub(r"\x01(\d+)\x01", lambda m: strings[int(m.group(1))], f)

    bare = _strip_strings(f)
    for fn in _FUNC.findall(bare):
        if fn.upper() in BANNED_FUNCS:
            problems.append(f"{fn.upper()} is not supported in Excel 2010/Google Sheets import; use SUMIFS/INDEX/MATCH")
    if "{" in bare or "}" in bare:
        problems.append(f"unsupported placeholder or array constant in {formula!r} (only {{r}}, {{r-1}}, {{r+1}})")
    depth = 0
    for ch in bare:
        depth += ch == "("
        depth -= ch == ")"
        if depth < 0:
            break
    if depth != 0:
        problems.append(f"unbalanced parentheses in {formula!r}")
    if not problems:
        try:
            Tokenizer(f)
        except Exception as e:  # noqa: BLE001
            problems.append(f"formula does not parse ({e}): {formula!r}")
    return f, problems


def guard_row_formula(f: str, row: int) -> str:
    """Blank on empty rows, never an error: =IF(COUNTA(inputs)=0,"",IFERROR(expr,""))."""
    body = f[1:]
    refs = []
    for col, r in _CELL.findall(_strip_strings(body)):
        ref = f"{col}{r}"
        if int(r) == row and ref not in refs:
            refs.append(ref)
    inner = f'IFERROR({body},"")'
    if refs:
        return f'=IF(COUNTA({",".join(refs[:30])})=0,"",{inner})'
    return "=" + inner


# ---------------------------------------------------------------------------------------------- values
def _parse_date(v: Any) -> Any:
    if isinstance(v, (date, datetime)):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%b %d, %Y", "%B %d, %Y", "%Y-%m"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return s


def _parse_number(v: Any, percent: bool = False) -> Any:
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return v / 100 if percent and abs(v) > 1.5 else v
    s = str(v).strip().replace(",", "").replace("$", "").replace("£", "").replace("€", "")
    is_pct = s.endswith("%")
    try:
        n = float(s.rstrip("%"))
    except ValueError:
        return str(v)
    return n / 100 if (is_pct or (percent and abs(n) > 1.5)) else n


def _cell_value(value: Any, col: dict[str, Any]) -> Any:
    if value is None or value == "":
        return None
    t = col["type"]
    if t == "date":
        return _parse_date(value)
    if t in ("currency", "number"):
        return _parse_number(value)
    if t == "percent":
        return _parse_number(value, percent=True)
    if t == "checkbox":
        return CHECK_ON if str(value).strip().lower() in ("true", "yes", "1", "x", CHECK_ON, "paid", "done") else CHECK_OFF
    return str(value)


# ---------------------------------------------------------------------------------------------- builder
class _Builder:
    def __init__(self, spec: dict[str, Any], palette: list[str] | None, snapshot: bool):
        palette = (palette or []) + ["#F7F3EC", "#2B2B2B", "#A3B18A"]
        self.spec = spec
        self.snapshot = snapshot
        self.bg, self.ink, self.accent = _hex(palette[0], "#F7F3EC"), _hex(palette[1], "#2B2B2B"), _hex(palette[2], "#A3B18A")
        if _luma(self.accent) > 0.82:  # too pale to carry white header text
            self.accent = _mix(self.accent, self.ink, 0.45)
        self.accent_dark = _mix(self.accent, "000000", 0.25)
        self.tint = _mix(self.accent, "FFFFFF", 0.86)
        self.tint2 = _mix(self.accent, "FFFFFF", 0.93)
        self.on_accent = "FFFFFF" if _luma(self.accent) < 0.6 else self.ink
        font = str(spec.get("font") or "Calibri")
        self.font = font if font in SAFE_FONTS else "Calibri"
        self.problems: list[str] = []
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.used: set[str] = {SETTINGS.lower(), START.lower()}
        self.sheet_map: dict[str, str] = {}  # lower(original or final) -> final name
        self.sheets: list[dict[str, Any]] = []
        self.lists: list[tuple[str, str, list[str]]] = []  # (defined name, header, options)
        self.thin = Side(style="thin", color="DADADA")
        self.border = Border(left=self.thin, right=self.thin, top=self.thin, bottom=self.thin)

    # -- small style helpers
    def f(self, **kw) -> Font:
        return Font(name=kw.pop("name", self.font), **kw)

    def fill(self, color: str) -> PatternFill:
        return PatternFill("solid", fgColor=color)

    def page_setup(self, ws, *, landscape: bool = True, fit_height: int = 0, title_rows: str | None = None) -> None:
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = fit_height
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.5
        ws.print_options.horizontalCentered = True
        if title_rows:
            ws.print_title_rows = title_rows

    # -- plan
    def plan_sheets(self) -> None:
        raw = [s for s in (self.spec.get("sheets") or []) if isinstance(s, dict)][:8]
        if not raw:
            raw = [{"name": "Sheet1", "columns": ["Date", "Item", "Amount"]}]
        for s in raw:
            original = str(s.get("name") or "Sheet")
            final = safe_sheet_name(original, self.used)
            self.sheet_map[original.strip().lower()] = final
            self.sheet_map[final.lower()] = final
            cols = normalize_columns(s.get("columns", []))
            rows = max(10, min(int(s.get("rows") or DEFAULT_ROWS), 500))
            self.sheets.append({"spec": s, "name": final, "cols": cols, "rows": rows,
                                "first": HEADER_ROW + 1, "last": HEADER_ROW + rows})
        dash = self.spec.get("dashboard")
        self.dash_name = None
        if isinstance(dash, dict) and (dash.get("kpis") or dash.get("breakdowns")):
            self.dash_name = safe_sheet_name(str(dash.get("name") or "Dashboard"), self.used)
            self.sheet_map[self.dash_name.lower()] = self.dash_name

    def find_sheet(self, name: Any) -> dict[str, Any] | None:
        final = self.sheet_map.get(str(name or "").strip().lower())
        return next((s for s in self.sheets if s["name"] == final), None)

    @staticmethod
    def find_col(sheet: dict[str, Any], name: Any) -> tuple[int, dict[str, Any]] | None:
        key = str(name or "").strip().lower()
        for i, c in enumerate(sheet["cols"], start=1):
            if c["name"].lower() == key:
                return i, c
        for i, c in enumerate(sheet["cols"], start=1):  # forgiving: partial match
            if key and (key in c["name"].lower() or c["name"].lower() in key):
                return i, c
        return None

    def col_range(self, sheet: dict[str, Any], idx: int) -> str:
        letter = get_column_letter(idx)
        return f"{_qs(sheet['name'])}!${letter}${sheet['first']}:${letter}${sheet['last']}"

    # -- tabs
    def build(self) -> Workbook:
        self.plan_sheets()
        start = self.wb.create_sheet(START)
        dash_ws = self.wb.create_sheet(self.dash_name) if self.dash_name else None
        for s in self.sheets:
            s["ws"] = self.wb.create_sheet(s["name"])
        for s in self.sheets:  # dropdown lists first: data tabs reference them
            for c in s["cols"]:
                if c["type"] == "choice":
                    c["list_name"] = f"list_{len(self.lists) + 1}"
                    self.lists.append((c["list_name"], f"{c['name']} ({s['name']})", c["options"]))
        if self.lists:  # a Settings tab only when there are dropdown lists to edit
            self.build_settings(self.wb.create_sheet(SETTINGS))
        for s in self.sheets:
            self.build_data_sheet(s)
        if dash_ws is not None:
            self.build_dashboard(dash_ws)
        self.build_start(start)
        self.wb.active = 0
        self.wb.calculation.fullCalcOnLoad = True
        return self.wb

    def banner(self, ws, title: str, subtitle: str, last_col: int) -> None:
        end = get_column_letter(max(2, last_col))
        ws.merge_cells(f"A1:{end}1")
        ws["A1"] = title
        ws["A1"].font = self.f(size=18, bold=True, color=self.on_accent)
        ws["A1"].alignment = Alignment(vertical="center", indent=1)
        for i in range(1, max(2, last_col) + 1):
            ws.cell(row=1, column=i).fill = self.fill(self.accent)
        ws.row_dimensions[1].height = 38
        ws.merge_cells(f"A2:{end}2")
        ws["A2"] = subtitle
        ws["A2"].font = self.f(size=10, italic=True, color="7A7A7A")
        ws["A2"].alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[2].height = 22

    def build_settings(self, ws) -> None:
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = "9E9E9E"
        n = max(1, len(self.lists))
        self.banner(ws, "Settings", "Edit these lists to rename or add your own dropdown options. "
                    "Changes flow through every tab and the dashboard.", n + 1)
        if not self.lists:
            ws["A4"] = "No dropdown lists in this template."
            ws["A4"].font = self.f(color="7A7A7A")
            return
        top = 4
        for i, (name, header, options) in enumerate(self.lists, start=1):
            letter = get_column_letter(i)
            h = ws.cell(row=top, column=i, value=header)
            h.font, h.fill, h.border = self.f(bold=True, color=self.on_accent), self.fill(self.accent), self.border
            h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[letter].width = max(18, min(34, len(header) + 4))
            for j in range(len(options) + SPARE_OPTIONS):
                c = ws.cell(row=top + 1 + j, column=i, value=options[j] if j < len(options) else None)
                c.border, c.font = self.border, self.f()
                if j >= len(options):
                    c.fill = self.fill(self.tint2)
            ref = f"{_qs(SETTINGS)}!${letter}${top + 1}:${letter}${top + len(options) + SPARE_OPTIONS}"
            self.wb.defined_names[name] = DefinedName(name, attr_text=ref)
            for c in self.sheets_cols_by_list(name):
                c["option_cells"] = [f"{_qs(SETTINGS)}!${letter}${top + 1 + j}" for j in range(len(options))]
        ws.row_dimensions[top].height = 32
        self.page_setup(ws, landscape=False)

    def sheets_cols_by_list(self, name: str):
        for s in self.sheets:
            for c in s["cols"]:
                if c.get("list_name") == name:
                    yield c

    def build_data_sheet(self, s: dict[str, Any]) -> None:
        ws, cols, first, last = s["ws"], s["cols"], s["first"], s["last"]
        spec = s["spec"]
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = self.accent
        purpose = str(spec.get("purpose") or "Type in the white cells; shaded columns calculate for you.")
        self.banner(ws, str(self.spec.get("title") or "Template"), f"{s['name']} · {purpose}", len(cols))
        samples = [r for r in (spec.get("sample_rows") or []) if isinstance(r, (list, tuple))][:12]
        if self.snapshot:
            last = min(last, first + max(len(samples), 3) + 6)
        checkbox_dv = None
        for i, col in enumerate(cols, start=1):
            letter = get_column_letter(i)
            disp = _display_type(col)
            head = ws.cell(row=HEADER_ROW, column=i, value=col["name"])
            head.fill, head.font, head.border = self.fill(self.accent), self.f(bold=True, color=self.on_accent), self.border
            head.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            width = col.get("width")
            try:
                width = int(width) if width else None
            except (TypeError, ValueError):
                width = None
            default_w = {"date": 15, "currency": 15, "number": 12, "percent": 12, "checkbox": 11, "choice": 18}.get(
                disp if col["type"] != "formula" else "currency", 22)
            ws.column_dimensions[letter].width = max(10, min(48, width or max(default_w, len(col["name"]) + 5)))
            formula_fill = self.fill(self.tint)
            for r in range(first, last + 1):
                c = ws.cell(row=r, column=i)
                c.border = self.border
                c.font = self.f()
                c.alignment = Alignment(vertical="center",
                                        horizontal="center" if col["type"] in ("checkbox", "date") else None,
                                        indent=1 if disp in ("text", "choice") else 0)
                if disp in NUMBER_FORMATS:
                    c.number_format = NUMBER_FORMATS[disp]
                if col["type"] == "formula":
                    c.fill = formula_fill
                    if col.get("formula"):
                        fx, problems = sanitize_formula(str(col["formula"]), row=r, sheet_map=self.sheet_map)
                        if problems:
                            if r == first:
                                self.problems += [f"{s['name']} › {col['name']}: {p}" for p in problems]
                            continue
                        c.value = guard_row_formula(fx, r)
                elif (r - first) % 2:
                    c.fill = self.fill("FAFAF8")
                k = r - first
                if col["type"] != "formula" and k < len(samples):
                    row_vals = list(samples[k]) + [None] * len(cols)
                    c.value = _cell_value(row_vals[i - 1], col)
            rng = f"{letter}{first}:{letter}{last}"
            if col["type"] == "checkbox":
                if checkbox_dv is None:
                    checkbox_dv = DataValidation(type="list", formula1=f'"{CHECK_OFF},{CHECK_ON}"', allow_blank=True)
                    ws.add_data_validation(checkbox_dv)
                checkbox_dv.add(rng)
                for r in range(first, last + 1):
                    ws.cell(row=r, column=i).font = self.f(size=13, color=self.accent_dark)
                ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[_str_lit(CHECK_ON)],
                                                              fill=self.fill("E3F1E4")))
            elif col["type"] == "choice":
                dv = DataValidation(type="list", formula1=f"={col['list_name']}", allow_blank=True,
                                    showErrorMessage=False)
                dv.promptTitle, dv.prompt = col["name"][:32], "Pick from the list (edit it on the Settings tab)"
                ws.add_data_validation(dv)
                dv.add(rng)
                if col.get("highlight") == "status" or any(w in col["name"].lower() for w in STATUS_WORDS):
                    for j, opt in enumerate(col["options"][:8]):
                        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[_str_lit(opt)],
                                                                      fill=self.fill(STATUS_TINTS[j % len(STATUS_TINTS)])))
            elif col["type"] == "date" and any(w in col["name"].lower() for w in ("due", "deadline")):
                ws.conditional_formatting.add(rng, FormulaRule(
                    formula=[f'AND({letter}{first}<>"",{letter}{first}<TODAY())'], font=Font(color="C0392B", bold=True)))
            if disp == "percent" and any(w in col["name"].lower() for w in ("progress", "complete", "goal", "saved", "paid off")):
                ws.conditional_formatting.add(rng, DataBarRule(start_type="num", start_value=0, end_type="num",
                                                               end_value=1, color=self.accent))
        totals = bool(spec.get("totals", True)) and not self.snapshot
        if totals:
            tr = last + 1
            for i, col in enumerate(cols, start=1):
                disp = _display_type(col)
                c = ws.cell(row=tr, column=i)
                c.fill, c.border = self.fill(self.tint), Border(top=Side(style="medium", color=self.accent), bottom=self.thin)
                if disp in SUMMABLE and col["type"] in ("currency", "number", "formula"):
                    letter = get_column_letter(i)
                    c.value = f"=SUM({letter}{first}:{letter}{last})"
                    c.font, c.number_format = self.f(bold=True), NUMBER_FORMATS[disp]
            label = ws.cell(row=tr, column=1)
            if label.value is None:
                label.value = "Total"
                label.font = self.f(bold=True)
            ws.row_dimensions[tr].height = 24
        ws.freeze_panes = f"A{first}"
        ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(cols))}{last}"
        ws.row_dimensions[HEADER_ROW].height = 30
        for r in range(first, last + 1):
            ws.row_dimensions[r].height = 20
        self.page_setup(ws, landscape=len(cols) > 4, fit_height=1 if self.snapshot else 0,
                        title_rows=f"{HEADER_ROW}:{HEADER_ROW}")
        if self.snapshot:
            ws.print_area = f"A1:{get_column_letter(len(cols))}{last}"

    # -- dashboard
    def kpi_formula(self, k: dict[str, Any], labels: dict[str, str]) -> tuple[str | None, str]:
        agg = str(k.get("agg") or "sum").lower()
        if agg in ("difference", "ratio"):
            of = [labels.get(str(x).strip().lower()) for x in (k.get("of") or [])][:2]
            if len(of) != 2 or None in of:
                return None, "number"
            a, b = of
            if agg == "difference":
                return f"={a}-{b}", "currency"
            return f'=IFERROR({a}/{b},"")', "percent"
        sheet = self.find_sheet(k.get("sheet"))
        if not sheet:
            return None, "number"
        found = self.find_col(sheet, k.get("column")) if k.get("column") else None
        where = k.get("where") if isinstance(k.get("where"), dict) else None
        cond = ""
        if where:
            wc = self.find_col(sheet, where.get("column"))
            if wc:
                eq = where.get("equals")
                crit = _str_lit(CHECK_ON) if wc[1]["type"] == "checkbox" and str(eq).lower() in ("true", "yes", "1", CHECK_ON) \
                    else _str_lit(CHECK_OFF) if wc[1]["type"] == "checkbox" else _str_lit(str(eq))
                cond = f"{self.col_range(sheet, wc[0])},{crit}"
        if agg == "count":
            target = found or (1, sheet["cols"][0])
            rng = self.col_range(sheet, target[0])
            return (f"=COUNTIFS({cond})" if cond else f"=SUMPRODUCT(--(LEN({rng})>0))"), "number"
        if not found:
            return None, "number"
        rng = self.col_range(sheet, found[0])
        fmt = _display_type(found[1])
        if agg == "average":
            return (f'=IFERROR(AVERAGEIFS({rng},{cond}),"")' if cond else f'=IFERROR(AVERAGE({rng}),"")'), fmt
        if agg in ("max", "min"):
            return f"={agg.upper()}({rng})", fmt
        return (f"=SUMIFS({rng},{cond})" if cond else f"=SUM({rng})"), fmt

    def build_dashboard(self, ws) -> None:
        dash = self.spec.get("dashboard") or {}
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = self.accent_dark
        for i in range(1, 11):
            ws.column_dimensions[get_column_letter(i)].width = 3 if i == 1 else 15
        ws.column_dimensions["J"].width = 3
        ws.merge_cells("B2:I2")
        ws["B2"] = str(dash.get("title") or self.spec.get("title") or "Dashboard")
        ws["B2"].font = self.f(size=22, bold=True, color=self.ink)
        ws.row_dimensions[2].height = 36
        ws.merge_cells("B3:F3")
        ws["B3"] = "Updates automatically as you fill in your tabs."
        ws["B3"].font = self.f(size=10, italic=True, color="7A7A7A")
        breakdowns = [b for b in (dash.get("breakdowns") or []) if isinstance(b, dict)][:3]
        year_cell = None
        if any(self._is_date_group(b) for b in breakdowns):
            ws["H3"], ws["I3"] = "Year", self._sample_year()
            ws["H3"].font = self.f(bold=True, color=self.accent_dark)
            ws["H3"].alignment = Alignment(horizontal="right")
            ws["I3"].font, ws["I3"].fill, ws["I3"].border = self.f(bold=True), self.fill("FFFFFF"), \
                Border(left=Side(style="medium", color=self.accent), right=Side(style="medium", color=self.accent),
                       top=Side(style="medium", color=self.accent), bottom=Side(style="medium", color=self.accent))
            ws["I3"].alignment = Alignment(horizontal="center")
            year_cell = "$I$3"
        # KPI cards: 4 per row, each two columns wide
        labels: dict[str, str] = {}
        kpis = [k for k in (dash.get("kpis") or []) if isinstance(k, dict)][:8]
        row = 5
        for n, k in enumerate(kpis):
            if n and n % 4 == 0:
                row += 4
            col = 2 + (n % 4) * 2
            formula, fmt = self.kpi_formula(k, labels)
            if formula is None:
                self.problems.append(f"Dashboard KPI '{k.get('label')}' references a sheet/column that does not exist")
                continue
            a, b = get_column_letter(col), get_column_letter(col + 1)
            for r in (row, row + 1, row + 2):
                for cc in (col, col + 1):
                    ws.cell(row=r, column=cc).fill = self.fill(self.tint)
            ws.merge_cells(f"{a}{row}:{b}{row}")
            ws.merge_cells(f"{a}{row + 1}:{b}{row + 1}")
            lc = ws[f"{a}{row}"]
            lc.value = str(k.get("label") or "Total").upper()
            lc.font = self.f(size=9, bold=True, color=self.accent_dark)
            lc.alignment = Alignment(horizontal="center", vertical="bottom")
            vc = ws[f"{a}{row + 1}"]
            vc.value = formula
            explicit = str(k.get("format") or "").lower()
            vc.number_format = NUMBER_FORMATS.get(explicit if explicit in NUMBER_FORMATS else fmt, "#,##0")
            if vc.number_format == NUMBER_FORMATS["number"]:
                vc.number_format = "#,##0"
            vc.font = self.f(size=20, bold=True, color=self.ink)
            vc.alignment = Alignment(horizontal="center", vertical="center")
            labels[str(k.get("label") or "").strip().lower()] = f"{a}{row + 1}"
            ws.row_dimensions[row + 1].height = 34
            ws.row_dimensions[row + 2].height = 8
        row += 4 if kpis else 0
        # Breakdown tables + native charts
        for b in breakdowns:
            row = self.build_breakdown(ws, b, row, year_cell) + 2
        self.page_setup(ws, landscape=False, fit_height=1)
        ws.print_area = f"A1:J{max(row, 20)}"

    def _is_date_group(self, b: dict[str, Any]) -> bool:
        sheet = self.find_sheet(b.get("sheet"))
        found = self.find_col(sheet, b.get("group_by")) if sheet else None
        return bool(found and found[1]["type"] == "date")

    def _sample_year(self) -> int:
        for s in self.sheets:
            for i, c in enumerate(s["cols"]):
                if c["type"] == "date":
                    for r in s["spec"].get("sample_rows") or []:
                        if isinstance(r, (list, tuple)) and i < len(r):
                            d = _parse_date(r[i])
                            if isinstance(d, (date, datetime)):
                                return d.year
        return date.today().year

    def build_breakdown(self, ws, b: dict[str, Any], row: int, year_cell: str | None) -> int:
        title = str(b.get("title") or "Breakdown")
        sheet = self.find_sheet(b.get("sheet"))
        group = self.find_col(sheet, b.get("group_by")) if sheet else None
        if not sheet or not group or group[1]["type"] not in ("choice", "date"):
            self.problems.append(f"Dashboard breakdown '{title}': group_by must be a dropdown (choice) or date column "
                                 f"on an existing sheet")
            return row
        agg = str(b.get("agg") or ("sum" if b.get("value_column") else "count")).lower()
        value = self.find_col(sheet, b.get("value_column")) if b.get("value_column") else None
        if agg == "sum" and not value:
            self.problems.append(f"Dashboard breakdown '{title}': value_column not found")
            return row
        grng = self.col_range(sheet, group[0])
        vrng = self.col_range(sheet, value[0]) if value else None
        fmt = _display_type(value[1]) if value and agg == "sum" else "number"
        ws.merge_cells(f"B{row}:I{row}")
        ws[f"B{row}"] = title
        ws[f"B{row}"].font = self.f(size=13, bold=True, color=self.ink)
        ws[f"B{row}"].border = Border(bottom=Side(style="medium", color=self.accent))
        ws.row_dimensions[row].height = 26
        head = row + 1
        for cc, text in ((2, group[1]["name"]), (3, value[1]["name"] if value and agg == "sum" else "Count")):
            h = ws.cell(row=head, column=cc, value=text)
            h.font, h.fill, h.border = self.f(bold=True, color=self.on_accent), self.fill(self.accent), self.border
            h.alignment = Alignment(horizontal="center")
        if group[1]["type"] == "date":
            labels: list[Any] = MONTHS
        else:
            labels = [f"={ref}" for ref in group[1].get("option_cells", [])]
        r = head
        for m, lab in enumerate(labels, start=1):
            r = head + m
            lc = ws.cell(row=r, column=2, value=lab)
            if group[1]["type"] == "date":
                lo = f'">="&DATE({year_cell},{m},1)'
                hi = f'"<"&DATE({year_cell},{m + 1},1)'
                crit = f"{grng},{lo},{grng},{hi}"
            else:
                crit = f"{grng},B{r}"
            formula = f"=SUMIFS({vrng},{crit})" if agg == "sum" else f"=COUNTIFS({crit})"
            vc = ws.cell(row=r, column=3, value=formula)
            vc.number_format = NUMBER_FORMATS.get(fmt, "#,##0") if fmt != "number" else "#,##0"
            for c in (lc, vc):
                c.border, c.font = self.border, self.f()
                if m % 2 == 0:
                    c.fill = self.fill("FAFAF8")
        kind = str(b.get("chart") or ("column" if group[1]["type"] == "date" else "doughnut")).lower()
        chart = self.make_chart(kind, ws, head, r, title)
        if chart is not None:
            ws.add_chart(chart, f"E{head}")
        return max(r, head + 15)

    def make_chart(self, kind: str, ws, head: int, last: int, title: str):
        data = Reference(ws, min_col=3, min_row=head, max_row=last)
        cats = Reference(ws, min_col=2, min_row=head + 1, max_row=last)
        if kind in ("pie", "doughnut"):
            chart = DoughnutChart(holeSize=55) if kind == "doughnut" else PieChart()
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            series = chart.series[0]
            shades = [_mix(self.accent, "FFFFFF", t) for t in (0, .25, .45, .6)] + \
                     [_mix(self.accent, self.ink, t) for t in (.3, .55)] + [_mix(self.accent, "000000", .45), "C9C3B8"]
            for i in range(last - head):
                pt = DataPoint(idx=i)
                pt.graphicalProperties.solidFill = shades[i % len(shades)]
                pt.graphicalProperties.line.solidFill = "FFFFFF"
                series.dPt.append(pt)
            chart.legend.position = "r"
        else:
            if kind == "line":
                chart = LineChart()
            else:
                chart = BarChart()
                chart.type = "bar" if kind == "bar" else "col"
                chart.gapWidth = 60
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            s = chart.series[0]
            s.graphicalProperties.solidFill = self.accent
            s.graphicalProperties.line.solidFill = self.accent
            if kind == "line":
                s.graphicalProperties.line.width = 28000
                s.smooth = False
            chart.legend = None
            chart.y_axis.majorGridlines = None
            chart.x_axis.delete = False
            chart.y_axis.delete = False
        chart.title = title
        chart.width, chart.height = 15.5, 7.6
        return chart

    # -- start here
    def build_start(self, ws) -> None:
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = self.ink
        ws.column_dimensions["A"].width = 3
        ws.column_dimensions["B"].width = 6
        ws.column_dimensions["C"].width = 30
        ws.column_dimensions["D"].width = 70
        for r in range(1, 4):
            for c in range(1, 6):
                ws.cell(row=r, column=c).fill = self.fill(self.accent)
        ws.merge_cells("B2:D2")
        ws["B2"] = str(self.spec.get("title") or "Template")
        ws["B2"].font = self.f(size=24, bold=True, color=self.on_accent)
        ws.row_dimensions[2].height = 44
        ws["B3"] = str(self.spec.get("subtitle") or "Your spreadsheet, ready to use")
        ws["B3"].font = self.f(size=11, italic=True, color=self.on_accent)
        ws.row_dimensions[3].height = 24
        row = 5

        def section(text: str) -> None:
            nonlocal row
            c = ws.cell(row=row, column=2, value=text.upper())
            c.font = self.f(size=11, bold=True, color=self.accent_dark)
            ws.merge_cells(f"B{row}:D{row}")
            c.border = Border(bottom=Side(style="thin", color=self.accent))
            ws.row_dimensions[row].height = 24
            row += 1

        def line(num: str, text: str, detail: str = "", *, bold: bool = False) -> None:
            nonlocal row
            ws.cell(row=row, column=2, value=num).font = self.f(bold=True, color=self.accent_dark)
            ws.cell(row=row, column=2).alignment = Alignment(horizontal="center", vertical="top")
            c = ws.cell(row=row, column=3, value=text)
            c.font = self.f(bold=bold)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if detail:
                d = ws.cell(row=row, column=4, value=detail)
                d.font, d.alignment = self.f(color="555555"), Alignment(wrap_text=True, vertical="top")
            else:
                ws.merge_cells(f"C{row}:D{row}")
            longest = max(len(text) / (95 if not detail else 28), len(detail) / 75)
            ws.row_dimensions[row].height = 18 * max(1, int(longest) + 1)
            row += 1

        steps = [str(x) for x in (self.spec.get("instructions") or []) if str(x).strip()][:8] or [
            "Start on the first data tab and replace the example rows with your own entries.",
            "Type only in white cells; shaded cells calculate automatically.",
            "Use the dropdowns to pick categories; customise the lists on the Settings tab." if self.lists
            else "Use the checkbox dropdowns to tick items off.",
            "Check the Dashboard for your totals and charts." if self.dash_name else "Totals update as you type."]
        section("How to use")
        for i, step in enumerate(steps, start=1):
            line(f"{i}", step)
        row += 1
        section("What's inside")
        if self.dash_name:
            line("•", self.dash_name, "Your totals and charts at a glance. Updates automatically.", bold=True)
        for s in self.sheets:
            line("•", s["name"], str(s["spec"].get("purpose") or "Log your entries here."), bold=True)
        if self.lists:
            line("•", SETTINGS, "Your dropdown lists (categories, statuses…). Rename or add options here.", bold=True)
        row += 1
        section("Colour key")
        keys = [("FFFFFF", "Type here", "White cells are yours to fill in."),
                (self.tint, "Calculated", "Shaded cells hold formulas; they fill in by themselves."),
                ("FFFFFF", "▾ Dropdown", "Click the cell and pick from the list.")]
        for color, label, detail in keys:
            c = ws.cell(row=row, column=3, value=label)
            c.fill, c.border, c.font = self.fill(color), self.border, self.f(bold=True)
            c.alignment = Alignment(horizontal="center")
            d = ws.cell(row=row, column=4, value=detail)
            d.font = self.f(color="555555")
            ws.row_dimensions[row].height = 20
            row += 1
        row += 1
        section("Good to know")
        tips = ["Example entries are included so you can see how everything works. Type over them, or select "
                "them and press Delete.",
                "Works in Microsoft Excel (2010 or newer) and Google Sheets (File › Import › Upload). "
                "Formulas and charts carry over.",
                "Add more rows any time by inserting rows inside a table, not below the totals row."]
        if self.dash_name and any(self._is_date_group(b) for b in (self.spec.get("dashboard") or {}).get("breakdowns") or []
                                  if isinstance(b, dict)):
            tips.append("On the Dashboard, change the Year box to see a different year.")
        for t in tips:
            line("✓", t)
        self.page_setup(ws, landscape=False, fit_height=1)


def render_spreadsheet(spec: dict[str, Any], *, palette: list[str] | None = None, snapshot: bool = False,
                       strict: bool = False) -> bytes:
    """Build the workbook. `snapshot` builds a print-tuned copy (one page per tab) for rendering images.
    `strict` raises SpreadsheetError when any formula/dashboard reference had to be dropped."""
    builder = _Builder(spec, palette, snapshot)
    wb = builder.build()
    if strict and builder.problems:
        raise SpreadsheetError(builder.problems)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def spec_problems(spec: dict[str, Any], palette: list[str] | None = None) -> list[str]:
    builder = _Builder(spec, palette, False)
    builder.build()
    return builder.problems


# ---------------------------------------------------------------------------------------------- validation
def validate_xlsx(data: bytes, *, recalc: bool = True) -> dict[str, Any]:
    """The gate every workbook passes before QA: opens cleanly, every formula parses, and (when
    LibreOffice is installed) a full recalculation produces zero error values."""
    errors: list[str] = []
    report: dict[str, Any] = {"ok": False, "errors": errors, "formulas": 0, "recalculated": False}
    try:
        wb = load_workbook(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001
        errors.append(f"workbook does not open: {e}")
        return report
    formulas = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    formulas += 1
                    bare = _strip_strings(v)
                    if "{" in bare or "}" in bare:
                        errors.append(f"{ws.title}!{cell.coordinate}: unreplaced placeholder {v[:60]}")
                    try:
                        Tokenizer(v)
                    except Exception as e:  # noqa: BLE001
                        errors.append(f"{ws.title}!{cell.coordinate}: formula does not parse ({e})")
    report["formulas"] = formulas
    report["sheets"] = wb.sheetnames
    if recalc:
        from pipeline.generators import office

        if office.available():
            try:
                calc = load_workbook(io.BytesIO(office.convert(data, "xlsx", "xlsx")), data_only=True)
                report["recalculated"] = True
                for ws in calc.worksheets:
                    for row in ws.iter_rows():
                        for cell in row:
                            if isinstance(cell.value, str) and cell.value.startswith(EXCEL_ERRORS):
                                errors.append(f"{ws.title}!{cell.coordinate} = {cell.value}")
            except Exception as e:  # noqa: BLE001
                errors.append(f"recalculation failed: {e}")
        else:
            report["note"] = "LibreOffice not installed; structural checks only"
    if len(errors) > 25:
        report["errors"] = errors[:25] + [f"... and {len(errors) - 25} more"]
    report["ok"] = not errors
    return report


def sheet_snapshots(spec: dict[str, Any], *, palette: list[str] | None = None) -> dict[str, bytes]:
    """Render each tab of the real workbook (LibreOffice -> PDF -> PNG). {} when LibreOffice is missing."""
    from pipeline.generators import office

    if not office.available():
        return {}
    data = render_spreadsheet(spec, palette=palette, snapshot=True)
    names = load_workbook(io.BytesIO(data), read_only=True).sheetnames
    pages = office.pdf_to_pngs(office.convert(data, "xlsx", "pdf"), scale=2.0, max_pages=len(names))
    return {name: office.trim(png) for name, png in zip(names, pages, strict=False)}


# ---------------------------------------------------------------------------------------------- previews
def _fmt(value: Any, ctype: str) -> str:
    if value is None or value == "":
        return ""
    if ctype == "currency":
        try:
            return f"${float(_parse_number(value)):,.2f}"
        except (TypeError, ValueError):
            return str(value)
    if ctype == "percent":
        try:
            v = float(_parse_number(value, percent=True))
            return f"{v * 100:.0f}%"
        except (TypeError, ValueError):
            return str(value)
    if ctype == "checkbox":
        return CHECK_ON if str(value).lower() in ("true", "yes", "1", "x", CHECK_ON, "paid", "done") else CHECK_OFF
    if ctype == "date":
        d = _parse_date(value)
        return d.strftime("%b %-d, %Y") if isinstance(d, (date, datetime)) else str(value)
    return str(value)


def window_html(image_png: bytes, *, title: str, subtitle: str, tabs: list[str], active: str,
                palette: list[str] | None = None, badge: str | None = None) -> str:
    """Frame a real rendered tab in a spreadsheet-app window for listing images."""
    import base64

    palette = (palette or []) + ["#F7F3EC", "#2B2B2B", "#A3B18A"]
    bg, ink, accent = palette[0], palette[1], palette[2]
    esc = html_lib.escape
    tab_html = "".join(f"<span class='tab{' on' if t == active else ''}'>{esc(t)}</span>" for t in tabs)
    b64 = base64.b64encode(image_png).decode()
    badge_html = f"<span class='badge'>{esc(badge)}</span>" if badge else ""
    return f"""<!doctype html><html><head><meta charset='utf-8'><style>
body{{margin:0;background:{bg};font-family:Inter,'DejaVu Sans',sans-serif;color:{ink}}}
.wrap{{padding:44px}} h1{{font-family:Poppins,Inter,sans-serif;font-size:40px;margin:0 0 6px}}
.sub{{opacity:.7;font-size:18px;margin-bottom:14px}}
.badge{{display:inline-block;margin-bottom:18px;padding:6px 12px;border-radius:999px;background:{accent}22;font-size:14px}}
.win{{background:#fff;border-radius:14px;box-shadow:0 18px 50px rgba(0,0,0,.16);overflow:hidden}}
.bar{{height:34px;background:#EEECE8;display:flex;align-items:center;gap:8px;padding:0 14px}}
.dot{{width:11px;height:11px;border-radius:50%;background:#D6D2CB}}
.tabs{{display:flex;gap:4px;padding:10px 14px 0;background:#F6F5F2;border-bottom:1px solid #E4E1DC;flex-wrap:wrap}}
.tab{{padding:7px 14px;font-size:13px;border-radius:8px 8px 0 0;background:#ECEAE6}}
.tab.on{{background:{accent};color:#fff;font-weight:600}}
.shot{{display:block;width:100%;height:auto}}
</style></head><body><div class='wrap'><h1>{esc(short_title(title))}</h1><div class='sub'>{esc(subtitle)}</div>{badge_html}
<div class='win'><div class='bar'><span class='dot'></span><span class='dot'></span><span class='dot'></span></div>
<div class='tabs'>{tab_html}</div><img class='shot' src='data:image/png;base64,{b64}'></div></div></body></html>"""


def preview_html(spec: dict[str, Any], *, palette: list[str] | None = None, sheets: list[int] | None = None,
                 subtitle: str | None = None) -> str:
    """A clean, spreadsheet-app-style preview (tabs, header row, example data) for listing images."""
    palette = palette or ["#F7F3EC", "#2B2B2B", "#A3B18A"]
    bg, ink, accent = (palette + ["#F7F3EC", "#2B2B2B", "#A3B18A"])[:3]
    esc = html_lib.escape
    all_sheets = spec.get("sheets") or []
    idx = [i for i in (sheets if sheets is not None else range(min(2, len(all_sheets)))) if i < len(all_sheets)]
    chosen = [all_sheets[i] for i in idx]
    active = idx[0] if idx else 0
    tabs = "".join(f"<span class='tab{' on' if i == active else ''}'>{esc(str(s.get('name', 'Sheet')))}</span>" for i, s in enumerate(all_sheets))
    widest = max((len(normalize_columns(s.get("columns", []))) for s in chosen), default=1)
    font = 14 if widest <= 6 else 12 if widest <= 8 else 11
    blocks = []
    for s in chosen:
        cols = normalize_columns(s.get("columns", []))[:9]
        head = "".join(f"<th class='{_display_type(c)}'>{esc(c['name'])}{' ▾' if c['type'] == 'choice' else ''}</th>" for c in cols)
        body = ""
        for row in (s.get("sample_rows") or [])[:6]:
            cells = list(row) + [""] * len(cols)
            body += "<tr>" + "".join(
                f"<td class='{_display_type(c)}{' calc' if c['type'] == 'formula' else ''}'>{esc(_fmt(cells[i], _display_type(c)))}</td>"
                for i, c in enumerate(cols)) + "</tr>"
        body += "<tr class='empty'>" + "<td></td>" * len(cols) + "</tr>"
        blocks.append(f"<div class='sheet'><div class='sname'>{esc(str(s.get('name', '')))}</div>"
                      f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>")
    return f"""<!doctype html><html><head><meta charset='utf-8'><style>
body{{margin:0;background:{bg};font-family:Inter,'DejaVu Sans',sans-serif;color:{ink}}}
.badge{{display:inline-block;margin-top:14px;padding:6px 12px;border-radius:999px;background:{accent}22;color:{ink};font-size:14px}}
.wrap{{padding:48px}} h1{{font-family:Poppins,Inter,sans-serif;font-size:40px;margin:0 0 6px}}
.sub{{opacity:.65;font-size:18px;margin-bottom:22px}}
.win{{background:#fff;border-radius:14px;box-shadow:0 18px 50px rgba(0,0,0,.14);overflow:hidden}}
.bar{{height:34px;background:#EEECE8;display:flex;align-items:center;gap:8px;padding:0 14px}}
.dot{{width:11px;height:11px;border-radius:50%;background:#D6D2CB}}
.tabs{{display:flex;gap:4px;padding:10px 14px 0;background:#F6F5F2;border-bottom:1px solid #E4E1DC;flex-wrap:wrap}}
.tab{{padding:7px 16px;font-size:14px;border-radius:8px 8px 0 0;background:#ECEAE6}}
.tab.on{{background:{accent};color:#fff;font-weight:600}}
.sheet{{padding:16px 18px 6px}} .sname{{font-weight:600;font-size:15px;margin:0 0 8px;color:{accent}}}
table{{border-collapse:collapse;width:100%;font-size:{font}px;table-layout:auto}}
th{{background:{accent};color:#fff;text-align:left;padding:9px 10px;font-weight:600}}
td{{padding:8px 10px;border-bottom:1px solid #EEE}} tr:nth-child(even) td{{background:#FAFAF8}}
td.calc{{background:{accent}18}}
td.currency,td.number,td.percent,td.formula,th.currency,th.number,th.percent,th.formula{{text-align:right;font-variant-numeric:tabular-nums}}
th.checkbox{{text-align:center}}
td.checkbox{{text-align:center;color:{accent};font-size:17px}} tr.empty td{{height:22px}}
</style></head><body><div class='wrap'><h1>{esc(short_title(str(spec.get('title', 'Template'))))}</h1>
<div class='sub'>{esc(subtitle or 'Editable spreadsheet · Excel & Google Sheets · formulas included')}</div>
<span class='badge'>Example entries included · type over them with your own</span><div style='height:18px'></div>
<div class='win'><div class='bar'><span class='dot'></span><span class='dot'></span><span class='dot'></span></div>
<div class='tabs'>{tabs}</div>{''.join(blocks)}</div></div></body></html>"""
